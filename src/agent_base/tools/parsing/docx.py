"""DOCX 解析器（python-docx，可选依赖 ``[doc]`` extras）。

python-docx 未安装时本模块仍可安全导入（守卫式导入）：``is_available()``
返回 False，注册表把 docx 视为未注册，``parse_document`` 给出指向
安装命令的可读错误——这正是"未安装对应 extras 时该格式解析器不注册"
的优雅降级契约。

只提取正文段落（与 chat-agent 的已验证行为一致）；表格、页眉页脚等
留到有真实需求时再扩展。空段落不进入输出（只为 LLM 省 token），
段落数元信息仍统计全部段落。
"""

from __future__ import annotations

import io

try:
    from docx import Document as _Document
except ImportError:  # 未安装 [doc] extras——按契约降级，不报错
    _Document = None  # type: ignore[assignment]

from agent_base.tools.parsing.base import DocumentParseError, ParsedDocument

DOCX_MISSING_MESSAGE = (
    "DOCX 解析不可用：未安装可选依赖 python-docx；请执行 pip install 'agent-base[doc]' 安装后重试"
)


def is_available() -> bool:
    """python-docx 是否可用；注册表据此决定 docx 是否注册。"""
    return _Document is not None


def parse_docx(data: bytes) -> ParsedDocument:
    """解析 DOCX 字节流，返回正文文本与总段落数。"""
    if _Document is None:
        raise DocumentParseError(DOCX_MISSING_MESSAGE)
    try:
        document = _Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentParseError(
            f"DOCX 解析失败：文件可能已损坏或不是有效的 .docx（{type(exc).__name__}: {exc}）"
        ) from exc
    paragraphs = list(document.paragraphs)
    text = "\n".join(p.text for p in paragraphs if p.text)
    return ParsedDocument(text=text, format="docx", paragraphs=len(paragraphs))
