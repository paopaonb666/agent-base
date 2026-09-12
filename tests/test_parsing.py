"""工具库 parsing 层的测试（M4 前置）。

覆盖任务规格中的验收项：各格式正常解析、``pages`` 页码选择、输入/
输出限额、损坏文件抛可读异常、``[doc]`` extras 缺失时的优雅降级。
PDF 夹具在测试内确定性构造（最小合法 PDF），不依赖二进制 fixture。
"""

from __future__ import annotations

import io

import pytest

from agent_base.tools.parsing import (
    DEFAULT_MAX_INPUT_BYTES,
    DEFAULT_MAX_OUTPUT_CHARS,
    DocumentParseError,
    ParsedDocument,
    available_formats,
    parse_document,
)
from agent_base.tools.parsing import docx as docx_module
from agent_base.tools.parsing.base import parse_page_selection

requires_docx = pytest.mark.skipif(
    not docx_module.is_available(), reason="python-docx 未安装（[doc] extras）"
)


# ── 夹具构造 ─────────────────────────────────────────────────────────


def _make_pdf(page_texts: list[str]) -> bytes:
    """构造最小但合法的多页 PDF：每页一行 Helvetica 文本。"""
    objects: list[bytes] = []
    count = len(page_texts)
    # 对象布局：1=catalog，2=pages 树，3=字体，之后每页两个对象（页面+内容流）。
    kids = " ".join(f"{4 + i * 2} 0 R" for i in range(count))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {count} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(page_texts):
        content_id = 5 + i * 2
        objects.append(
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>".encode()
        )
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def _make_encrypted_pdf(page_texts: list[str]) -> bytes:
    """构造带密码的 PDF（pypdf 自身的加密能力）。"""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.append(io.BytesIO(_make_pdf(page_texts)))
    writer.encrypt("secret")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _make_docx(paragraphs: list[str]) -> bytes:
    """构造内存中的 .docx（python-docx 自身的写入能力）。"""
    import docx as python_docx

    document = python_docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ── PDF ──────────────────────────────────────────────────────────────


def test_pdf_parses_all_pages_by_default() -> None:
    result = parse_document(_make_pdf(["Hello one", "Hello two"]), filename="doc.pdf")
    assert result.format == "pdf"
    assert result.pages == 2
    assert result.truncated is False
    assert result.text == "Hello one\nHello two"


def test_pdf_pages_single_page_selection() -> None:
    result = parse_document(_make_pdf(["one", "two", "three"]), filename="doc.pdf", pages="2")
    assert result.text == "two"
    # pages 元信息始终是文档总页数，供上层提示"还有更多页"。
    assert result.pages == 3


def test_pdf_pages_range_and_mixed() -> None:
    result = parse_document(
        _make_pdf(["p1", "p2", "p3", "p4", "p5"]), filename="d.pdf", pages="1-3,5"
    )
    assert result.text == "p1\np2\np3\np5"


def test_pdf_pages_dedup_and_sorted() -> None:
    # 重复与乱序的页码选择按阅读顺序去重输出，不会重复提取。
    result = parse_document(_make_pdf(["p1", "p2", "p3"]), filename="d.pdf", pages="3,1,1-2")
    assert result.text == "p1\np2\np3"


def test_pdf_pages_out_of_range_is_readable_error() -> None:
    with pytest.raises(DocumentParseError, match="共 3 页"):
        parse_document(_make_pdf(["p1", "p2", "p3"]), filename="d.pdf", pages="1,4")


@pytest.mark.parametrize(
    ("spec", "reason"),
    [
        ("abc", "不是整数页码"),
        ("3-2", "小于起点"),
        ("0", "页码从 1 开始"),
        ("1,,2", "空的页码项"),
        ("1-2-3", "不是页码或页码区间"),
        ("-2", "不是整数页码"),
        ("", "空的页码项"),
    ],
)
def test_pdf_pages_malformed(spec: str, reason: str) -> None:
    with pytest.raises(DocumentParseError, match=reason):
        parse_document(_make_pdf(["p1", "p2"]), filename="d.pdf", pages=spec)


def test_page_selection_direct() -> None:
    assert parse_page_selection("1-3,5", total=5) == [1, 2, 3, 5]
    assert parse_page_selection("2-2", total=2) == [2]


def test_pdf_encrypted_rejected_with_readable_error() -> None:
    with pytest.raises(DocumentParseError, match="已加密"):
        parse_document(_make_encrypted_pdf(["secret"]), filename="d.pdf")


def test_pdf_corrupted_raises_not_silent() -> None:
    with pytest.raises(DocumentParseError, match="PDF 解析失败"):
        parse_document(b"%PDF-1.4 this is not really a pdf", filename="d.pdf")


def test_pdf_scanned_pages_yield_empty_text_not_error() -> None:
    # 扫描版（纯图片页）提取为空文本是合法结果，保留页数元信息。
    result = parse_document(_make_pdf(["", ""]), filename="scan.pdf")
    assert result.text.strip() == ""
    assert result.pages == 2


# ── TXT / Markdown ───────────────────────────────────────────────────


def test_txt_utf8_decode_and_paragraphs() -> None:
    result = parse_document("第一行\nsecond line\n\n".encode(), filename="a.txt")
    assert result.format == "txt"
    assert result.text == "第一行\nsecond line\n\n"
    # 段落数按非空行统计；空行不计。
    assert result.paragraphs == 2


def test_txt_gbk_fallback() -> None:
    result = parse_document("中文内容，GBK 编码。".encode("gbk"), filename="a.txt")
    assert result.text == "中文内容，GBK 编码。"


def test_txt_undecodable_raises() -> None:
    # 0x80 在 UTF-8 与 GB18030 中都不是合法字节——必须报错而非静默替换。
    with pytest.raises(DocumentParseError, match="解码失败"):
        parse_document(b"\x80\x80\x80", filename="a.txt")


def test_txt_utf8_bom_is_stripped() -> None:
    result = parse_document(b"\xef\xbb\xbfhello", filename="a.txt")
    assert result.text == "hello"


def test_empty_txt_is_valid() -> None:
    result = parse_document(b"", filename="empty.txt")
    assert result.text == ""
    assert result.paragraphs == 0
    assert result.truncated is False


@pytest.mark.parametrize("filename", ["notes.md", "notes.markdown"])
def test_markdown_format_normalized(filename: str) -> None:
    result = parse_document("# 标题\n\n正文".encode(), filename=filename)
    assert result.format == "markdown"
    assert result.paragraphs == 2


def test_pages_parameter_ignored_for_text_formats() -> None:
    # pages 仅 PDF 支持；文本格式显式忽略（契约写在 parse_document docstring）。
    result = parse_document(b"alpha\nbeta", filename="a.txt", pages="1-3")
    assert result.text == "alpha\nbeta"


# ── DOCX（可选 extras） ──────────────────────────────────────────────


@requires_docx
def test_docx_roundtrip() -> None:
    result = parse_document(_make_docx(["First", "", "Second"]), filename="r.docx")
    assert result.format == "docx"
    # 空段落不进入输出文本（省 token），但计入段落元信息。
    assert result.text == "First\nSecond"
    assert result.paragraphs == 3


@requires_docx
def test_docx_corrupted_raises_not_silent() -> None:
    with pytest.raises(DocumentParseError, match="DOCX 解析失败"):
        parse_document(b"PK\x03\x04 not a real docx", filename="r.docx")


def test_docx_unavailable_degrades_gracefully(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """模拟未安装 [doc] extras：docx 不注册，报错给出安装指引而非崩溃。"""
    monkeypatch.setattr(docx_module, "_Document", None)
    assert "docx" not in available_formats()
    with pytest.raises(DocumentParseError, match=r"agent-base\[doc\]"):
        parse_document(b"whatever", filename="r.docx")
    # 绕过注册表直接调用解析器同样得到可读错误（守卫分支兜底）。
    with pytest.raises(DocumentParseError, match="python-docx"):
        docx_module.parse_docx(b"whatever")


# ── 分发 / 注册表 ────────────────────────────────────────────────────


def test_available_formats_contains_baselines() -> None:
    assert {"pdf", "txt", "markdown"} <= set(available_formats())


def test_unknown_extension_readable_error() -> None:
    with pytest.raises(DocumentParseError, match="无法识别的文档扩展名 'doc'"):
        parse_document(b"...", filename="legacy.doc")


def test_missing_filename_and_format_raises() -> None:
    with pytest.raises(DocumentParseError, match="无法确定文档格式"):
        parse_document(b"...")


def test_explicit_format_overrides_suffix() -> None:
    result = parse_document(b"plain text", filename="report.pdf", format="txt")
    assert result.format == "txt"


def test_suffix_case_insensitive() -> None:
    result = parse_document(_make_pdf(["x"]), filename="REPORT.PDF")
    assert result.format == "pdf"


def test_result_is_plain_dataclass() -> None:
    # 解析结果是纯数据（无 IO 句柄），可被工具层与 RAG 模块安全持有。
    result = parse_document(b"hello", filename="a.txt")
    assert isinstance(result, ParsedDocument)
    assert result == ParsedDocument(text="hello", format="txt", paragraphs=1)


# ── 限额 ─────────────────────────────────────────────────────────────


def test_output_truncation_sets_flag() -> None:
    result = parse_document(b"a" * 100, filename="a.txt", max_output_chars=10)
    assert result.text == "a" * 10
    assert result.truncated is True


def test_output_under_limit_not_flagged() -> None:
    result = parse_document(b"a" * 10, filename="a.txt", max_output_chars=10)
    assert result.truncated is False


def test_input_size_cap_rejects() -> None:
    with pytest.raises(DocumentParseError, match="超出大小上限"):
        parse_document(b"x" * 100, filename="a.txt", max_input_bytes=8)


def test_nonpositive_limits_rejected() -> None:
    with pytest.raises(DocumentParseError, match="必须为正数"):
        parse_document(b"x", filename="a.txt", max_output_chars=0)
    with pytest.raises(DocumentParseError, match="必须为正数"):
        parse_document(b"x", filename="a.txt", max_input_bytes=-1)


def test_default_limits_match_config_contract() -> None:
    # parsing 层的默认限额与 core/config.py 的契约字段必须一致，
    # 否则"走配置"和"走默认值"两条路径会出现静默不同的行为。
    from agent_base.core.config import Settings

    settings = Settings(_env_file=None)
    assert settings.doc_parse_max_input_bytes == DEFAULT_MAX_INPUT_BYTES
    assert settings.doc_parse_max_output_chars == DEFAULT_MAX_OUTPUT_CHARS
